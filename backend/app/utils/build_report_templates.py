"""Builds the Word (.docx) report templates used by report_service.

Run from the `backend/` directory:

    python -m app.utils.build_report_templates

This produces `Default_Audit_Report.docx` plus one copy per framework listed
in report_service.TEMPLATE_MAP, so every registered framework works
immediately. To customise a framework's look, open its .docx in Microsoft
Word and edit the layout - just keep the Jinja2 tags ({{ ... }} / {% ... %})
intact. (Re-running this script will overwrite those edits, so only re-run it
if you want to reset templates to the generated baseline.)

Template table rules (docxtpl):
  * `{%tr for row in controls %}` and `{%tr endfor %}` must each sit in their
    own table row. docxtpl repeats the row between them and removes the
    tag rows.
"""

import os

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

from app.services.report_service import DEFAULT_TEMPLATE, TEMPLATE_MAP, TEMPLATES_DIR

BRAND = RGBColor(0x0F, 0x17, 0x2A)

# (header text, jinja field on each control row, column width in cm)
COLUMNS = [
    ("S.No", "sno", 1.2),
    ("Clause (Framework Rules)", "clause", 2.6),
    ("Description of Finding(s)", "description", 5.5),
    ("System (RE / Vendor)", "system", 2.2),
    ("Status", "status", 2.2),
    ("Risk (C/H/M/L)", "risk", 1.6),
    ("C/I/A", "cia", 1.4),
    ("Test Cases Used", "test_cases", 2.6),
    ("Auditor Notes", "auditor_notes", 3.2),
    ("Root Cause Analysis", "root_cause", 2.6),
    ("Impact Analysis", "impact", 2.6),
    ("Recommendations / Corrective Actions", "recommendations", 3.2),
    ("Days", "days", 1.2),
    ("Management Response", "mgmt_response", 2.6),
    ("Similar Issue (Last 3 Audits)", "similar_issue", 2.6),
    ("Evidence Required", "evidence", 3.0),
]


def _shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _set_cell(cell, text: str, *, bold=False, size=8, color=None, center=False):
    cell.text = ""
    para = cell.paragraphs[0]
    run = para.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if center:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _heading(doc, text, size=13):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = BRAND
    return para


def build_template(path: str) -> None:
    doc = Document()

    # Landscape - the controls table is 16 columns wide.
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(1.2))

    # ── Title block ──────────────────────────────────────────────
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("{{ audit_type }} Audit Report")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = BRAND

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("{{ company_name }}").font.size = Pt(14)

    # ── Audit details ────────────────────────────────────────────
    _heading(doc, "1. Audit Details")
    details = [
        ("Company Name", "{{ company_name }}"),
        ("Registration No.", "{{ registration_no }}"),
        ("Audit Name", "{{ audit_name }}"),
        ("Framework", "{{ audit_type }}"),
        ("Category", "{{ audit_category }}"),
        ("Sub-category", "{{ audit_subcategory }}"),
        ("Financial Year", "{{ target_fy }}"),
        ("Audit Status", "{{ audit_status }}"),
        ("Report Generated On", "{{ generated_date }}"),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in details:
        row = table.add_row()
        _set_cell(row.cells[0], label, bold=True, size=10)
        _shade(row.cells[0], "E2E8F0")
        _set_cell(row.cells[1], value, size=10)
        row.cells[0].width = Cm(6)
        row.cells[1].width = Cm(12)

    # ── Executive summary ────────────────────────────────────────
    doc.add_paragraph()
    _heading(doc, "2. Executive Summary")
    doc.add_paragraph("Total controls assessed: {{ total_controls }}")

    summary = doc.add_table(rows=3, cols=2)
    summary.style = "Table Grid"
    _set_cell(summary.rows[0].cells[0], "Status", bold=True, size=10)
    _set_cell(summary.rows[0].cells[1], "Count", bold=True, size=10)
    _shade(summary.rows[0].cells[0], "E2E8F0")
    _shade(summary.rows[0].cells[1], "E2E8F0")
    _set_cell(summary.rows[1].cells[0], "{%tr for s in status_counts %}", size=10)
    _set_cell(summary.rows[1].cells[1], "", size=10)
    _set_cell(summary.rows[2].cells[0], "{{ s.status }}", size=10)
    _set_cell(summary.rows[2].cells[1], "{{ s.count }}", size=10)
    end_row = summary.add_row()
    _set_cell(end_row.cells[0], "{%tr endfor %}", size=10)
    _set_cell(end_row.cells[1], "", size=10)

    # ── Detailed findings table ──────────────────────────────────
    doc.add_paragraph()
    _heading(doc, "3. Detailed Audit Findings")

    ctrl_table = doc.add_table(rows=1, cols=len(COLUMNS))
    ctrl_table.style = "Table Grid"
    ctrl_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, (header, _field, width) in enumerate(COLUMNS):
        cell = ctrl_table.rows[0].cells[i]
        _set_cell(cell, header, bold=True, size=8, color=RGBColor(0xFF, 0xFF, 0xFF))
        _shade(cell, "0F172A")
        cell.width = Cm(width)

    loop_row = ctrl_table.add_row()
    _set_cell(loop_row.cells[0], "{%tr for row in controls %}", size=8)

    data_row = ctrl_table.add_row()
    for i, (_h, field, width) in enumerate(COLUMNS):
        _set_cell(data_row.cells[i], "{{ row." + field + " }}", size=8)
        data_row.cells[i].width = Cm(width)

    end_loop_row = ctrl_table.add_row()
    _set_cell(end_loop_row.cells[0], "{%tr endfor %}", size=8)

    # ── Sign-off ─────────────────────────────────────────────────
    doc.add_paragraph()
    _heading(doc, "4. Auditor Sign-off")
    sign = doc.add_table(rows=2, cols=3)
    sign.style = "Table Grid"
    for i, label in enumerate(("Lead Auditor", "Signature", "Date")):
        _set_cell(sign.rows[0].cells[i], label, bold=True, size=10)
        _shade(sign.rows[0].cells[i], "E2E8F0")
        _set_cell(sign.rows[1].cells[i], "\n\n", size=10)

    os.makedirs(os.path.dirname(path), exist_ok=True)
    doc.save(path)


def main() -> None:
    filenames = {DEFAULT_TEMPLATE, *TEMPLATE_MAP.values()}
    for name in sorted(filenames):
        target = os.path.normpath(os.path.join(TEMPLATES_DIR, name))
        build_template(target)
        print(f"built {target}")


if __name__ == "__main__":
    main()